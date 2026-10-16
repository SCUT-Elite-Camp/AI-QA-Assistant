"""Native XLSX parser that preserves worksheets and merged cells."""

from __future__ import annotations

from openpyxl import load_workbook

from models.document import BlockType, ContentBlock, Document, TableArtifact, TableCell
from parsers.base import BaseParser


class XlsxParser(BaseParser):
    def parse(self, file_path: str) -> Document:
        workbook = load_workbook(file_path, data_only=False, read_only=False)
        cached = load_workbook(file_path, data_only=True, read_only=False)
        blocks: list[ContentBlock] = []
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                block = self._sheet_to_block(sheet, sheet_index, cached[sheet.title])
                if block is not None:
                    blocks.append(block)
        finally:
            workbook.close()
            cached.close()

        content = "\n\n".join(block.to_markdown() for block in blocks if not block.is_empty)
        return Document.from_file_path(file_path, content, content_blocks=blocks)

    @staticmethod
    def _sheet_to_block(sheet, sheet_index: int, cached_sheet=None) -> ContentBlock | None:
        nonempty = [
            cell
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        ]
        if not nonempty:
            return None

        min_row = min(cell.row for cell in nonempty)
        max_row = max(cell.row for cell in nonempty)
        min_column = min(cell.column for cell in nonempty)
        max_column = max(cell.column for cell in nonempty)

        # Merged placeholders have no value. A populated merged heading can
        # extend beyond the last nonempty cell, so include its full geometry.
        retained_merges = [
            merged for merged in sheet.merged_cells.ranges
            if not (merged.max_row < min_row or merged.min_row > max_row
                    or merged.max_col < min_column or merged.min_col > max_column)
        ]
        for merged in retained_merges:
            min_row = min(min_row, merged.min_row)
            max_row = max(max_row, merged.max_row)
            min_column = min(min_column, merged.min_col)
            max_column = max(max_column, merged.max_col)

        merged_starts: dict[tuple[int, int], tuple[int, int]] = {}
        merged_covered: set[tuple[int, int]] = set()
        for merged in retained_merges:
            start = (merged.min_row, merged.min_col)
            merged_starts[start] = (
                merged.max_row - merged.min_row + 1,
                merged.max_col - merged.min_col + 1,
            )
            for row_index in range(merged.min_row, merged.max_row + 1):
                for column_index in range(merged.min_col, merged.max_col + 1):
                    if (row_index, column_index) != start:
                        merged_covered.add((row_index, column_index))

        cells: list[TableCell] = []
        bold_by_row: list[tuple[int, int]] = []
        for source_row in range(min_row, max_row + 1):
            bold = 0
            populated = 0
            for source_column in range(min_column, max_column + 1):
                if (source_row, source_column) in merged_covered:
                    continue
                source = sheet.cell(source_row, source_column)
                value = "" if source.value is None else str(source.value).strip()
                formula = value if source.data_type == 'f' else None
                cached_value = None
                if formula:
                    cached_value = cached_sheet.cell(source_row, source_column).value if cached_sheet is not None else None
                    value = str(cached_value) if cached_value is not None else f'{formula} [result unavailable]'
                row_span, column_span = merged_starts.get(
                    (source_row, source_column),
                    (1, 1),
                )
                cells.append(TableCell(
                    row_index=source_row - min_row,
                    column_index=source_column - min_column,
                    text=value,
                    formula=formula,
                    cached_value=str(cached_value) if cached_value is not None else None,
                    row_span=row_span,
                    column_span=column_span,
                ))
                if value:
                    populated += 1
                    if source.font and source.font.bold:
                        bold += 1
            bold_by_row.append((bold, populated))

        header_row_count = _xlsx_header_rows(cells, bold_by_row, max_column - min_column + 1)
        for cell in cells:
            cell.is_header = cell.row_index < header_row_count

        artifact = TableArtifact(
            table_id=f"xlsx-sheet-{sheet_index}",
            cells=cells,
            row_count=max_row - min_row + 1,
            column_count=max_column - min_column + 1,
            header_row_count=header_row_count,
            caption=sheet.title,
            parser_name="openpyxl",
            confidence=0.99,
            backend_confidence=0.99,
            quality_score=1.0,
            quality_evidence={"native_structure": 1.0},
            acceptance="accepted",
        )
        from openpyxl.utils import get_column_letter
        return ContentBlock(
            block_type=BlockType.TABLE, table=artifact,
            confidence=0.7 if any(cell.data_type == 'f' for cell in nonempty) else 1.0,
            locator={"sheet": sheet.title, "cell_range": f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"},
        )

    @staticmethod
    def supported_extensions() -> list[str]:
        return [".xlsx"]


def _xlsx_header_rows(
    cells: list[TableCell],
    bold_by_row: list[tuple[int, int]],
    column_count: int,
) -> int:
    header_count = 0
    for bold, populated in bold_by_row:
        if populated and bold / populated >= 0.5:
            header_count += 1
        else:
            break
    if header_count:
        return header_count

    matrix = [["" for _ in range(column_count)] for _ in range(len(bold_by_row))]
    for cell in cells:
        for row_index in range(cell.row_index, cell.row_index + cell.row_span):
            for column_index in range(cell.column_index, cell.column_index + cell.column_span):
                matrix[row_index][column_index] = cell.text
    if len(matrix) < 2:
        return 0
    # Native spreadsheets often use two or more text-only header rows followed
    # by numeric data (for example, a merged year above Q1/Q2 columns).
    for row_index, row in enumerate(matrix[1:], start=1):
        values = [value for value in row if value]
        if values and any(_is_number(value) for value in values):
            return row_index
    return 0


def _is_number(value: str) -> bool:
    try:
        float(value.replace(",", "").replace("%", ""))
    except ValueError:
        return False
    return True
