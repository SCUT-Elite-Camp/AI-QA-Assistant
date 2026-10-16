import zipfile
from openpyxl import Workbook
from models.document import ContentBlock, Document
from parsers.xlsx_parser import XlsxParser
from parsers.table_backends import _native_cell_order
from pipeline.ingestion.preparation import prepare_document
from pipeline.documents.budget import count_tokens, split_budget, token_limit


def test_budget_split_is_lossless_and_does_not_discard_tail():
    text = ''.join(f'Step {i} requires a separate check before proceeding. ' for i in range(180))
    parts = split_budget(text)
    assert ''.join(parts) == text
    assert len(parts) > 1
    assert all(count_tokens(part) <= token_limit() for part in parts)


def test_long_document_keeps_paths_and_explicit_neighbors():
    text = ''.join(f'Requirement {i}: preserve the original source. ' for i in range(100))
    doc = Document(doc_id='long', title='Design', content=text, space='test', address='long.md', last_updated='2026-10-16', content_blocks=[
        ContentBlock(block_type='heading', text='Requirements', level=1),
        ContentBlock(block_type='paragraph', text=text)])
    prepare_document(doc)
    assert len(doc.chunks) > 1
    assert all(c.section_path == ['Requirements'] for c in doc.chunks)
    assert all(count_tokens(c.text) <= token_limit() for c in doc.chunks)
    assert all('Design' in c.index_text for c in doc.chunks)
    assert doc.chunks[0].metadata['next_chunk_id'] == doc.chunks[1].chunk_id
    assert doc.chunks[-1].metadata['previous_chunk_id'] == doc.chunks[-2].chunk_id


def test_large_table_preserves_every_row_binding_without_full_table_copies():
    rows = [[f'TEST-{i:03}', f'{i*13}', f'Check result for owner {i}'] for i in range(90)]
    doc = Document(doc_id='table', title='Test report', content='', space='test', address='table.md', last_updated='2026-10-16', content_blocks=[
        ContentBlock(block_type='table', headers=['Test ID', 'Duration (ms)', 'Result'], rows=rows)])
    prepare_document(doc)
    assert 1 < len(doc.chunks) < 90
    assert len({c.text for c in doc.chunks}) == len(doc.chunks)
    for row in rows:
        matching = [c for c in doc.chunks if row[0] in c.text]
        assert len(matching) == 1
        assert row[1] in matching[0].text
        assert 'Duration (ms)' in matching[0].text
    assert all(count_tokens(c.text) <= token_limit() for c in doc.chunks)
    assert all(c.metadata['row_start'] <= c.metadata['row_end'] for c in doc.chunks)


def test_xlsx_keeps_formula_and_existing_cached_result(tmp_path):
    path = tmp_path/'cached.xlsx'
    wb = Workbook(); sheet = wb.active
    sheet.append(['Item', 'Total']);sheet.append(['Alpha', '=2+3']);sheet.append(['Beta', '=4+5'])
    wb.save(path);wb.close()
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    sheet_xml = entries['xl/worksheets/sheet1.xml'].decode()
    assert '<f>2+3</f><v></v>' in sheet_xml
    entries['xl/worksheets/sheet1.xml'] = sheet_xml.replace('<f>2+3</f><v></v>', '<f>2+3</f><v>5</v>').encode()
    with zipfile.ZipFile(path, 'w') as archive:
        for name, blob in entries.items():archive.writestr(name, blob)
    table = XlsxParser().parse(str(path)).content_blocks[0].table
    alpha = next(c for c in table.cells if c.formula == '=2+3')
    beta = next(c for c in table.cells if c.formula == '=4+5')
    assert alpha.text == '5' and alpha.cached_value == '5'
    assert beta.cached_value is None and 'result unavailable' in beta.text


def test_native_cell_order_preserves_exact_naming_characters():
    class Row:
        cells = [(0, 0, 100, 20)]
    class Table:
        rows = [Row()]
        def extract(self):return [['YYYY-MM-DD <name>\n_']]
    class Page:
        def get_text(self, *args, **kwargs):return 'YYYY-MM-DD_<name>'
    assert _native_cell_order(Page(), Table()) == [['YYYY-MM-DD_<name>']]


def test_clipped_text_cannot_add_or_drop_characters():
    class Row:
        cells = [(0, 0, 100, 20)]
    class Table:
        rows = [Row()]
        def extract(self):return [['Original 12']]
    class Page:
        def get_text(self, *args, **kwargs):return 'Other 120'
    assert _native_cell_order(Page(), Table()) == [['Original 12']]


def test_native_draw_order_does_not_reorder_multiline_numeric_data():
    class Row:
        cells = [(0,0,100,20)]
    class Table:
        rows = [Row()]
        def extract(self):return [['12\n34']]
    class Page:
        def get_text(self, *args, **kwargs):return '34\n12'
    assert _native_cell_order(Page(), Table()) == [['12\n34']]


def test_header_recovery_stops_before_data_instead_of_promoting_later_prose():
    doc = Document(doc_id='metrics', title='Metrics', content='', space='test', address='metrics.md',
        last_updated='2026-10-16', content_blocks=[ContentBlock(block_type='table', rows=[
            ['Total tests', '34', 'Collected items'], ['Pass rate', '', 'Passed / total'],
            ['Unit tests', '31', 'Service and retrieval']])])
    prepare_document(doc)
    assert not doc.chunks[0].metadata['header_inferred']
    assert 'Column 1' in doc.chunks[0].text
    assert 'Pass rate' in doc.chunks[0].text


def test_table_groups_repeat_explicit_units_with_their_source_block():
    doc = Document(doc_id='units', title='Revenue', content='', space='test', address='units.md',
        last_updated='2026-10-16', content_blocks=[
            ContentBlock(block_type='paragraph', text='All amounts are in thousands of USD.', block_id='units'),
            ContentBlock(block_type='table', headers=['Region','Revenue'],
                         rows=[[f'Region-{i}',str(i*10)] for i in range(70)]),
            ContentBlock(block_type='paragraph', text='Note: provisional figures.', block_id='note')])
    prepare_document(doc)
    tables = [c for c in doc.chunks if c.chunk_type == 'table']
    assert len(tables) > 1
    assert all('thousands of USD' in c.text and 'provisional figures' in c.text for c in tables)
    assert all(c.metadata['context_block_ids'] == ['units', 'note'] for c in tables)
    assert all(count_tokens(c.text) <= token_limit() for c in tables)
