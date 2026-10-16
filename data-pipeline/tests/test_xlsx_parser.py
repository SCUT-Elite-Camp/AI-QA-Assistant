import sys
from pathlib import Path

from openpyxl import Workbook


PIPELINE_ROOT = Path(__file__).resolve().parents[1]
if str(PIPELINE_ROOT) not in sys.path:
    sys.path.insert(0, str(PIPELINE_ROOT))

from models.document import BlockType  # noqa: E402
from parsers.registry import DocumentParser  # noqa: E402
from parsers.xlsx_parser import XlsxParser  # noqa: E402


def test_xlsx_parser_preserves_merged_cells_and_sheet_provenance(tmp_path):
    source = tmp_path / "sales.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["地区", "销售额", None])
    sheet.append([None, "Q1", "Q2"])
    sheet.append(["华南", 120, 136])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:C1")
    workbook.save(source)

    parsed = XlsxParser().parse(str(source))

    assert len(parsed.content_blocks) == 1
    block = parsed.content_blocks[0]
    assert block.block_type == BlockType.TABLE
    assert block.table.caption == "销售"
    assert block.table.parser_name == "openpyxl"
    assert block.table.legacy_headers() == ["地区", "销售额 / Q1", "销售额 / Q2"]
    assert block.table.body_rows() == [["华南", "120", "136"]]


def test_default_registry_supports_xlsx():
    parser = DocumentParser()

    assert ".xlsx" in parser.supported_extensions
    assert isinstance(parser.get_parser(".xlsx"), XlsxParser)


def test_merged_heading_extends_past_last_populated_row_and_column(tmp_path):
    source = tmp_path / "merged-edge.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet['B2'] = 'wide heading'
    sheet.merge_cells('B2:F2')
    sheet['B3'] = 'tall note'
    sheet.merge_cells('B3:B7')
    sheet['C3'] = 20
    workbook.save(source)
    workbook.close()

    table = XlsxParser().parse(str(source)).content_blocks[0].table

    assert (table.row_count, table.column_count) == (6, 5)
    heading = next(cell for cell in table.cells if cell.text == 'wide heading')
    note = next(cell for cell in table.cells if cell.text == 'tall note')
    assert heading.column_span == 5
    assert note.row_span == 5
    assert table.body_rows()  # Rendering also respects the retained bounds.
