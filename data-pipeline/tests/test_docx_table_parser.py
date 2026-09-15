import sys
from pathlib import Path

from docx import Document as DocxDocument


PIPELINE_ROOT = Path(__file__).resolve().parents[1]
if str(PIPELINE_ROOT) not in sys.path:
    sys.path.insert(0, str(PIPELINE_ROOT))

from models.document import BlockType  # noqa: E402
from parsers.docx_parser import DocxParser  # noqa: E402


def set_bold(cell):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True


def test_docx_parser_preserves_horizontal_and_vertical_merge_semantics(tmp_path):
    source = tmp_path / "merged.docx"
    doc = DocxDocument()
    doc.add_paragraph("before table")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Sales"
    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(1, 1).text = "Q1"
    table.cell(1, 2).text = "Q2"
    table.cell(2, 0).text = "South"
    table.cell(2, 1).text = "120"
    table.cell(2, 2).text = "136"
    for row in table.rows[:2]:
        for cell in row.cells:
            set_bold(cell)
    doc.add_paragraph("after table")
    doc.save(source)

    parsed = DocxParser().parse(str(source))

    assert [block.block_type for block in parsed.content_blocks] == [
        BlockType.PARAGRAPH,
        BlockType.TABLE,
        BlockType.PARAGRAPH,
    ]
    artifact = parsed.content_blocks[1].table
    assert artifact is not None
    assert artifact.header_row_count == 2
    region = next(cell for cell in artifact.cells if cell.text == "Region")
    sales = next(cell for cell in artifact.cells if cell.text == "Sales")
    assert region.row_span == 2
    assert region.column_span == 1
    assert sales.row_span == 1
    assert sales.column_span == 2
    assert artifact.legacy_headers() == ["Region", "Sales / Q1", "Sales / Q2"]
    assert artifact.body_rows() == [["South", "120", "136"]]
    assert artifact.acceptance == "accepted"
    assert artifact.quality_evidence == {"native_structure": 1.0}


def test_docx_parser_does_not_assume_unstyled_first_row_is_header(tmp_path):
    source = tmp_path / "plain.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Alice"
    table.cell(0, 1).text = "Engineering"
    table.cell(1, 0).text = "Bob"
    table.cell(1, 1).text = "Finance"
    doc.save(source)

    parsed = DocxParser().parse(str(source))

    block = parsed.content_blocks[0]
    assert block.table.header_row_count == 0
    assert block.headers == []
    assert block.rows == [["Alice", "Engineering"], ["Bob", "Finance"]]
