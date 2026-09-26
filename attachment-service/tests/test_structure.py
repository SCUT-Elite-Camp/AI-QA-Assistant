from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation

from attachment_service.structure import build_document_sections


def _attachment(identifier: str, extension: str, filename: str) -> dict:
    return {
        "id": identifier,
        "version_id": identifier,
        "extension": extension,
        "filename": filename,
    }


def test_builds_markdown_heading_tree_from_chunk_locators(tmp_path: Path):
    path = tmp_path / "policy.md"
    path.write_text("# Scope\nA\n## Controls\nB", encoding="utf-8")
    evidence = [
        {"evidence_id": "e1", "content": "# Scope\nA", "locator": {"section_path": ["Scope"]}},
        {"evidence_id": "e2", "content": "## Controls\nB", "locator": {"section_path": ["Scope", "Controls"]}},
    ]
    rows = build_document_sections(path, _attachment("v1", ".md", "policy.md"), evidence)
    assert [row["title"] for row in rows] == ["policy.md", "Scope", "Controls"]
    assert rows[-1]["evidence_ids"] == ["e2"]
    assert rows[-1]["parent_id"] == rows[1]["id"]


def test_builds_pdf_toc_ranges(tmp_path: Path):
    path = tmp_path / "manual.pdf"
    with fitz.open() as document:
        document.new_page().insert_text((72, 72), "Introduction")
        document.new_page().insert_text((72, 72), "Controls")
        document.set_toc([[1, "Introduction", 1], [1, "Controls", 2]])
        document.save(path)
    evidence = [
        {"evidence_id": "p1", "content": "Introduction", "locator": {"page": 1}},
        {"evidence_id": "p2", "content": "Controls", "locator": {"page": 2}},
    ]
    rows = build_document_sections(path, _attachment("v2", ".pdf", "manual.pdf"), evidence)
    assert rows[1]["page_start"] == rows[1]["page_end"] == 1
    assert rows[2]["evidence_ids"] == ["p2"]


def test_builds_docx_headings_and_pptx_slides(tmp_path: Path):
    docx_path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Scope", level=1)
    document.add_paragraph("Scope text")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "table"
    document.add_heading("Controls", level=2)
    document.add_paragraph("Control text")
    document.save(docx_path)
    docx_evidence = [
        {"evidence_id": "d1", "content": "Scope", "locator": {"block": 0}},
        {"evidence_id": "d2", "content": "Controls", "locator": {"block": 3}},
    ]
    docx_rows = build_document_sections(
        docx_path, _attachment("v3", ".docx", "policy.docx"), docx_evidence,
    )
    assert [row["title"] for row in docx_rows] == ["policy.docx", "Scope", "Controls"]
    assert docx_rows[-1]["evidence_ids"] == ["d2"]

    pptx_path = tmp_path / "deck.pptx"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = "Overview"
    second = presentation.slides.add_slide(presentation.slide_layouts[0])
    second.shapes.title.text = "Risks"
    presentation.save(pptx_path)
    pptx_evidence = [
        {"evidence_id": "s1", "content": "Overview", "locator": {"slide": 1}},
        {"evidence_id": "s2", "content": "Risks", "locator": {"slide": 2}},
    ]
    pptx_rows = build_document_sections(
        pptx_path, _attachment("v4", ".pptx", "deck.pptx"), pptx_evidence,
    )
    assert [row["title"] for row in pptx_rows] == ["deck.pptx", "Overview", "Risks"]
    assert pptx_rows[2]["evidence_ids"] == ["s2"]
