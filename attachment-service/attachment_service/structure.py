from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Iterable, Protocol


class DocumentStructureProvider(Protocol):
    def build(
        self,
        path: Path,
        attachment: dict[str, Any],
        evidence: list[dict[str, Any]],
    ) -> list[dict[str, Any]]: ...


def _section_id(version_id: str, path: Iterable[str]) -> str:
    value = " / ".join(path)
    digest = hashlib.sha256(f"{version_id}:{value}".encode("utf-8")).hexdigest()[:20]
    return f"sec_{digest}"


def _summary(items: list[dict[str, Any]], limit: int = 800) -> str:
    text = "\n".join(str(item.get("content") or "").strip() for item in items).strip()
    return text[:limit]


def _row(
    version_id: str,
    attachment_id: str,
    path: list[str],
    evidence: list[dict[str, Any]],
    *,
    parent_path: list[str] | None = None,
    page_start: int | None = None,
    page_end: int | None = None,
    quality: str = "high",
    ordinal: int = 0,
) -> dict[str, Any]:
    return {
        "id": _section_id(version_id, path),
        "attachment_id": attachment_id,
        "version_id": version_id,
        "parent_id": _section_id(version_id, parent_path) if parent_path else None,
        "level": max(0, len(path) - 1),
        "title": path[-1],
        "section_path": path,
        "page_start": page_start,
        "page_end": page_end,
        "summary": _summary(evidence),
        "evidence_ids": [str(item["evidence_id"]) for item in evidence],
        "quality": quality,
        "ordinal": ordinal,
    }


def _markdown_sections(
    attachment_id: str,
    version_id: str,
    filename: str,
    evidence: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    root_path = [filename]
    rows = [_row(version_id, attachment_id, root_path, evidence, quality="high")]
    paths: list[list[str]] = []
    for item in evidence:
        raw = (item.get("locator") or {}).get("section_path") or []
        if isinstance(raw, list) and raw:
            values = [str(value).strip() for value in raw if str(value).strip()]
            for depth in range(1, len(values) + 1):
                path = root_path + values[:depth]
                if path not in paths:
                    paths.append(path)
    for ordinal, path in enumerate(paths, 1):
        members = [
            item for item in evidence
            if (
                root_path + list((item.get("locator") or {}).get("section_path") or [])
            )[:len(path)] == path
        ]
        rows.append(_row(
            version_id, attachment_id, path, members,
            parent_path=path[:-1], ordinal=ordinal,
        ))
    return rows


def _pdf_sections(
    path: Path,
    attachment_id: str,
    version_id: str,
    filename: str,
    evidence: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    import fitz

    with fitz.open(path) as document:
        toc = document.get_toc(simple=True)
        page_count = document.page_count
    root_path = [filename]
    rows = [_row(
        version_id, attachment_id, root_path, evidence,
        page_start=1 if page_count else None, page_end=page_count or None,
        quality="high" if toc else "low",
    )]
    if not toc:
        return rows
    stack: list[tuple[int, str]] = []
    for ordinal, (level, title, page) in enumerate(toc, 1):
        title = re.sub(r"\s+", " ", str(title)).strip()
        if not title:
            continue
        stack = [entry for entry in stack if entry[0] < int(level)]
        stack.append((int(level), title))
        next_pages = [int(item[2]) for item in toc[ordinal:] if int(item[0]) <= int(level)]
        page_end = (next_pages[0] - 1) if next_pages else page_count
        members = [
            item for item in evidence
            if int((item.get("locator") or {}).get("page") or 0) in range(int(page), page_end + 1)
        ]
        section_path = root_path + [entry[1] for entry in stack]
        rows.append(_row(
            version_id, attachment_id, section_path, members,
            parent_path=section_path[:-1], page_start=int(page), page_end=page_end,
            ordinal=ordinal,
        ))
    return rows


def _docx_sections(
    path: Path,
    attachment_id: str,
    version_id: str,
    filename: str,
    evidence: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    from docx import Document

    root_path = [filename]
    rows = [_row(version_id, attachment_id, root_path, evidence, quality="high")]
    headings: list[tuple[int, str, int]] = []
    # Keep locator numbering identical to the DOCX evidence parser: paragraphs
    # and tables both consume one position in the document body.
    for block, item in enumerate(Document(path).iter_inner_content()):
        if not (hasattr(item, "style") and hasattr(item, "text")):
            continue
        text = item.text.strip()
        style = str(item.style.name if item.style else "")
        match = re.search(r"([1-6])", style)
        normalized_style = style.lower()
        is_heading = (
            normalized_style.startswith("heading")
            or "标题" in normalized_style
            or normalized_style.startswith("title")
        )
        if text and is_heading:
            headings.append((int(match.group(1)) if match else 1, text, block))
    stack: list[tuple[int, str]] = []
    for ordinal, (level, title, start) in enumerate(headings, 1):
        stack = [entry for entry in stack if entry[0] < level]
        stack.append((level, title))
        end = headings[ordinal][2] - 1 if ordinal < len(headings) else 2**31 - 1
        members = [
            item for item in evidence
            if start <= int((item.get("locator") or {}).get("block") or 0) <= end
        ]
        section_path = root_path + [entry[1] for entry in stack]
        rows.append(_row(
            version_id, attachment_id, section_path, members,
            parent_path=section_path[:-1], ordinal=ordinal,
        ))
    if not headings:
        rows[0]["quality"] = "low"
    return rows


def _pptx_sections(
    path: Path,
    attachment_id: str,
    version_id: str,
    filename: str,
    evidence: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    from pptx import Presentation

    presentation = Presentation(path)
    root_path = [filename]
    rows = [_row(version_id, attachment_id, root_path, evidence, quality="high")]
    for slide_number, slide in enumerate(presentation.slides, 1):
        title = ""
        if slide.shapes.title is not None:
            title = str(slide.shapes.title.text or "").strip()
        title = title or f"Slide {slide_number}"
        members = [
            item for item in evidence
            if int((item.get("locator") or {}).get("slide") or 0) == slide_number
        ]
        section_path = root_path + [title]
        rows.append(_row(
            version_id, attachment_id, section_path, members,
            parent_path=root_path, page_start=slide_number, page_end=slide_number,
            ordinal=slide_number,
        ))
    return rows


class NativeStructureProvider:
    """Native P0 provider for Markdown, PDF, DOCX, and PPTX."""

    def build(
        self,
        path: Path,
        attachment: dict[str, Any],
        evidence: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        attachment_id = str(attachment["id"])
        version_id = str(attachment.get("version_id") or attachment_id)
        filename = str(attachment.get("filename") or attachment_id)
        extension = str(attachment.get("extension") or path.suffix).lower()
        if extension == ".md":
            return _markdown_sections(attachment_id, version_id, filename, evidence)
        if extension == ".pdf":
            return _pdf_sections(path, attachment_id, version_id, filename, evidence)
        if extension == ".docx":
            return _docx_sections(path, attachment_id, version_id, filename, evidence)
        if extension == ".pptx":
            return _pptx_sections(path, attachment_id, version_id, filename, evidence)
        return [_row(version_id, attachment_id, [filename], evidence, quality="low")]


def build_document_sections(
    path: Path,
    attachment: dict[str, Any],
    evidence: list[dict[str, Any]],
    provider: DocumentStructureProvider | None = None,
) -> list[dict[str, Any]]:
    """Build version-derived navigation rows without changing Evidence authority."""
    return (provider or NativeStructureProvider()).build(path, attachment, evidence)
